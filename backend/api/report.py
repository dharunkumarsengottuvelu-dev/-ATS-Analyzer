from fastapi import APIRouter, Depends, HTTPException, BackgroundTasks
from fastapi.responses import FileResponse
from pydantic import BaseModel
from sqlalchemy.orm import Session
from backend.database.session import get_db, SessionLocal
from backend.schemas.analysis import MatchMetrics
from backend.scoring.engine import calculate_ats_score
from backend.ats.llm import get_resume_review
from backend.reports.generator import generate_pdf_report
from backend.models.report import Report
from backend.models.analysis import Analysis
from backend.models.user import User
from backend.api.deps import get_current_user_optional
from typing import Optional
import os
import uuid
import time
import logging

logger = logging.getLogger(__name__)

router = APIRouter(
    prefix="/report",
    tags=["Report"]
)

# In-memory dictionary to track job statuses
jobs = {}

class ReportRequest(BaseModel):
    filename: str
    resume_text: str
    job_description: str
    resume_data: dict
    metrics: MatchMetrics
    analysis_id: Optional[int] = None

async def _generate_report_task(job_id: str, request: ReportRequest, user_id: int):
    try:
        start_time = time.time()
        jobs[job_id] = {"status": "processing", "message": "Analyzing with LLM..."}
        logger.info(f"Job {job_id}: Started ATS scoring")
        
        # 1. ATS Rule Engine Score
        ats_score = calculate_ats_score(request.metrics, request.resume_data)
        
        # 2. Local LLM Review
        logger.info(f"Job {job_id}: Running local LLM inference...")
        llm_feedback = await get_resume_review(request.resume_text, request.job_description)
        
        # 3. Generate Reports (PDF, JSON, CSV)
        logger.info(f"Job {job_id}: Generating Reports...")
        jobs[job_id]["message"] = "Generating PDF, JSON, and CSV Documents..."
        pdf_path = generate_pdf_report(request.filename, ats_score, request.metrics, llm_feedback)
        
        # Generate JSON
        import json
        json_path = pdf_path.replace(".pdf", ".json")
        with open(json_path, "w", encoding="utf-8") as f:
            json.dump({
                "filename": request.filename,
                "ats_score": ats_score,
                "metrics": request.metrics.dict(),
                "feedback": llm_feedback
            }, f, indent=4)
            
        # Generate DOCX
        from docx import Document
        docx_path = pdf_path.replace(".pdf", ".docx")
        doc = Document()
        doc.add_heading(f"ATS Analysis Report: {request.filename}", 0)
        doc.add_heading(f"Score: {ats_score}/100", 1)
        doc.add_heading("Feedback", 2)
        doc.add_paragraph(llm_feedback)
        doc.save(docx_path)
            
        # Generate CSV
        import csv
        csv_path = pdf_path.replace(".pdf", ".csv")
        with open(csv_path, "w", encoding="utf-8", newline="") as f:
            writer = csv.writer(f)
            writer.writerow(["Filename", "ATS Score", "Missing Skills", "Matched Skills", "Feedback"])
            writer.writerow([
                request.filename, 
                ats_score, 
                ", ".join(request.metrics.missing_keywords),
                ", ".join(request.metrics.matched_keywords),
                llm_feedback
            ])
        
        # 4. Save to Database
        db = SessionLocal()
        try:
            db_report = Report(
                user_id=user_id,
                analysis_id=request.analysis_id if request.analysis_id else 1,
                format="multiple",
                file_path=pdf_path,
                llm_feedback=llm_feedback
            )
            db.add(db_report)
            
            if request.analysis_id:
                db_analysis = db.query(Analysis).filter(Analysis.id == request.analysis_id).first()
                if db_analysis:
                    db_analysis.ats_score = ats_score
                    
            db.commit()
        finally:
            db.close()
            
        elapsed = time.time() - start_time
        logger.info(f"Job {job_id}: Completed in {elapsed:.2f}s")
        
        jobs[job_id] = {
            "status": "completed",
            "message": "Report generated successfully",
            "pdf_url": f"/api/v1/report/download?filepath={pdf_path}",
            "json_url": f"/api/v1/report/download?filepath={json_path}",
            "csv_url": f"/api/v1/report/download?filepath={csv_path}",
            "ats_score": ats_score,
            "llm_feedback": llm_feedback,
            "execution_time": round(elapsed, 2)
        }
        
    except Exception as e:
        logger.error(f"Job {job_id}: Failed with error: {str(e)}")
        jobs[job_id] = {
            "status": "failed",
            "message": f"Report generation failed: {str(e)}"
        }

@router.post("/generate")
async def generate_report(
    request: ReportRequest, 
    background_tasks: BackgroundTasks,
    current_user: User = Depends(get_current_user_optional)
):
    """
    Initiates report generation in the background to avoid HTTP timeouts.
    """
    job_id = str(uuid.uuid4())
    jobs[job_id] = {"status": "pending", "message": "Job queued..."}
    
    user_id = current_user.id if current_user else 1
    background_tasks.add_task(_generate_report_task, job_id, request, user_id)
    
    return {
        "status": "success",
        "message": "Report generation job started",
        "job_id": job_id
    }

@router.get("/status/{job_id}")
async def get_report_status(job_id: str):
    if job_id not in jobs:
        raise HTTPException(status_code=404, detail="Job not found")
    return jobs[job_id]

@router.get("/download")
async def download_report(analysis_id: int, format: str = "pdf", db: Session = Depends(get_db)):
    report = db.query(Report).filter(Report.analysis_id == analysis_id).first()
    if not report:
        raise HTTPException(status_code=404, detail="Report not found")
        
    filepath = report.file_path
    if format == "json":
        filepath = filepath.replace(".pdf", ".json")
    elif format == "csv":
        filepath = filepath.replace(".pdf", ".csv")
    elif format == "docx":
        filepath = filepath.replace(".pdf", ".docx")
        
    if not os.path.exists(filepath):
        raise HTTPException(status_code=404, detail=f"Report format {format} not found on disk")
        
    ext = filepath.split(".")[-1].lower()
    media_types = {
        "pdf": "application/pdf",
        "json": "application/json",
        "csv": "text/csv",
        "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    }
    return FileResponse(path=filepath, filename=os.path.basename(filepath), media_type=media_types.get(ext, "application/octet-stream"))

from backend.models.resume import Resume
from backend.models.job import JobDescription
from sqlalchemy import desc

class ReportItem(BaseModel):
    id: int
    resume_filename: str
    job_title: str
    ats_score: float
    created_at: str
    
    class Config:
        from_attributes = True

@router.get("/", response_model=list[ReportItem])
def get_reports(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user_optional),
    skip: int = 0,
    limit: int = 50
):
    user_id = current_user.id if current_user else 1
    reports = db.query(Report).filter(Report.user_id == user_id).order_by(desc(Report.created_at)).offset(skip).limit(limit).all()
    
    results = []
    for r in reports:
        analysis = db.query(Analysis).filter(Analysis.id == r.analysis_id).first()
        if not analysis:
            continue
        resume = db.query(Resume).filter(Resume.id == analysis.resume_id).first()
        job = db.query(JobDescription).filter(JobDescription.id == analysis.job_id).first()
        
        results.append(ReportItem(
            id=r.id,
            resume_filename=resume.filename if resume else "Unknown",
            job_title=job.title if job else "Job Description",
            ats_score=analysis.ats_score,
            created_at=r.created_at.isoformat() if r.created_at else ""
        ))
    return results

@router.get("/{report_id}")
def get_report_detail(
    report_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user_optional)
):
    user_id = current_user.id if current_user else 1
    report = db.query(Report).filter(Report.id == report_id, Report.user_id == user_id).first()
    if not report:
        raise HTTPException(status_code=404, detail="Report not found")
        
    analysis = db.query(Analysis).filter(Analysis.id == report.analysis_id).first()
    resume = db.query(Resume).filter(Resume.id == analysis.resume_id).first() if analysis else None
    
    return {
        "id": report.id,
        "resume_filename": resume.filename if resume else "Unknown",
        "ats_score": analysis.ats_score if analysis else 0,
        "llm_feedback": report.llm_feedback,
        "file_path": report.file_path,
        "created_at": report.created_at
    }

@router.get("/{report_id}/download")
async def download_report_by_id(
    report_id: int, 
    format: str = "pdf", 
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user_optional)
):
    user_id = current_user.id if current_user else 1
    report = db.query(Report).filter(Report.id == report_id, Report.user_id == user_id).first()
    if not report:
        raise HTTPException(status_code=404, detail="Report not found")
        
    filepath = report.file_path
    if format == "json":
        filepath = filepath.replace(".pdf", ".json")
    elif format == "csv":
        filepath = filepath.replace(".pdf", ".csv")
    elif format == "docx":
        filepath = filepath.replace(".pdf", ".docx")
        
    if not os.path.exists(filepath):
        raise HTTPException(status_code=404, detail=f"Report format {format} not found on disk")
        
    ext = filepath.split(".")[-1].lower()
    media_types = {
        "pdf": "application/pdf",
        "json": "application/json",
        "csv": "text/csv",
        "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    }
    return FileResponse(path=filepath, filename=os.path.basename(filepath), media_type=media_types.get(ext, "application/octet-stream"))

@router.delete("/{report_id}")
def delete_report(
    report_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user_optional)
):
    user_id = current_user.id if current_user else 1
    report = db.query(Report).filter(Report.id == report_id, Report.user_id == user_id).first()
    
    if not report:
        raise HTTPException(status_code=404, detail="Report not found")
        
    # Optional: Delete file from disk
    if report.file_path and os.path.exists(report.file_path):
        try:
            os.remove(report.file_path)
        except:
            pass
            
    db.delete(report)
    db.commit()
    return {"status": "success", "message": "Report deleted"}
